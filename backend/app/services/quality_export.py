"""Download of peer-reviewed sources: one workbook with the peers' comments beside the data, so the same file can be
handed to an AI model (or a person) to double-check what the peers said."""
import io
import json
from datetime import timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

KIND = {"comment": "Peer comment", "mention": "Peer mentioned", "finding": "Error logged", "sop": "SOP", "system": "System check", "ai": "AI or web flag"}
README = [
    "Peer-reviewed sources: what the reviewers said, next to the data.",
    "",
    "Sheets",
    "Sources: one row per source with counts.",
    "Peer comments: every comment a reviewer wrote (source notes and notes on a single field or record).",
    "Fields: one row per field of every record. Doubted fields come first. 'Why doubted' lists every signal, 'Peer says' quotes the reviewers.",
    "Records JSON: the raw record as JSON, so a model can read a whole record in one cell.",
    "",
    "How to use it with an AI model",
    "1. Give the model the Fields sheet (or the JSON download) and the source website.",
    "2. Ask it to check each row where 'Doubted' is yes against the website and answer correct, wrong or unsure, with the real value.",
    "3. Compare its answers with the 'Admin verdict' column. Anything it calls wrong that no peer pointed at is worth a look.",
    "A signal is only a pointer. It is not proof that a value is wrong.",
]


def _fmt(dt):
    if not dt:
        return ""
    try:
        return (dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)).strftime("%Y-%m-%d %H:%M")
    except Exception:
        return str(dt)


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    return v


def collect(items: list[dict]) -> dict:
    """items: [{source, project, extractor, review, peer_rows, sop_code}] -> plain structure used by both formats."""
    sources, comments, fields, records = [], [], [], []
    for it in items:
        s, rv = it["source"], it["review"]
        tot = rv["totals"]
        sources.append({"project": it["project"], "source": s.name, "website": s.website_url or "", "sop": rv.get("sop_code") or "",
                        "extractor": it["extractor"], "peers": len(rv["peers"]), "records": len(rv["records"]), "fields": tot["fields"], "doubted": tot["flagged"],
                        "peer_comments": tot["comments"], "peer_mentions": tot["mentions"], "checked": tot["audited"]})
        for p in it["peer_rows"]:
            comments.append({"project": it["project"], "source": s.name, "record": "(whole source)", "field": "", "by": p["name"], "role": "peer reviewer", "type": p["status"], "comment": p["comment"], "when": p["when"]})
        for r in rv["records"]:
            for g in r["general"]:
                comments.append({"project": it["project"], "source": s.name, "record": r["record"], "field": "", "by": g["by"], "role": g.get("role") or "", "type": g.get("type") or "note", "comment": g["text"], "when": g.get("ts") or ""})
            raw = {}
            for f in r["fields"]:
                raw[f["name"]] = f["value"]
                peer_say = " | ".join(f"{x.get('by') or 'Reviewer'}: {x['text']}" for x in f["signals"] if x["kind"] == "comment")
                for x in f["signals"]:
                    if x["kind"] == "comment":
                        comments.append({"project": it["project"], "source": s.name, "record": r["record"], "field": f["name"], "by": x.get("by") or "", "role": x.get("role") or "", "type": "field comment", "comment": x["text"], "when": ""})
                mention = " | ".join(f"{x.get('by') or 'Reviewer'}: {x['text']}" for x in f["signals"] if x["kind"] == "mention")
                why = "; ".join(sorted({KIND.get(x["kind"], x["kind"]) for x in f["signals"]}))
                a = f["audit"] or {}
                fields.append({"project": it["project"], "source": s.name, "website": s.website_url or "", "record": r["record"], "field": f["name"], "value": f["value"],
                               "doubted": "yes" if f["signals"] else "no", "why": why, "peer_says": peer_say, "peer_mentioned_in": mention,
                               "other": "; ".join(x["text"] for x in f["signals"] if x["kind"] in ("sop", "system", "ai", "finding") and x.get("text"))[:500],
                               "verdict": {"correct": "correct", "incorrect": "wrong"}.get(a.get("verdict"), ""), "truth": a.get("truth_value") or ""})
            records.append({"project": it["project"], "source": s.name, "record": r["record"], "json": json.dumps(raw, ensure_ascii=False)})
    fields.sort(key=lambda x: (x["doubted"] != "yes", x["project"], x["source"], x["record"]))
    return {"sources": sources, "comments": comments, "fields": fields, "records": records}


def to_json(data: dict, meta: dict) -> bytes:
    return json.dumps({"about": README[0], "meta": meta, **data}, ensure_ascii=False, indent=1, default=str).encode("utf-8")


def to_xlsx(data: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    for line in README:
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 120

    def sheet(title, headers, rows):
        sh = wb.create_sheet(title)
        sh.append([h for h, _ in headers])
        for r in rows:
            sh.append([_cell(r.get(k)) for _, k in headers])
        for c in sh[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="2563EB")
        sh.freeze_panes = "A2"
        sh.auto_filter.ref = sh.dimensions
        for i, (h, k) in enumerate(headers, 1):
            longest = max([len(str(h))] + [len(str(_cell(r.get(k)))) for r in rows[:200]])
            sh.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 10), 60)
            for row in sh.iter_rows(min_row=2, min_col=i, max_col=i):
                for c in row:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
                    if isinstance(c.value, str) and c.value.startswith("="):
                        c.data_type = "s"
        return sh

    sheet("Sources", [("Project", "project"), ("Source", "source"), ("Website", "website"), ("SOP", "sop"), ("Extractor", "extractor"), ("Peers", "peers"), ("Records", "records"),
                      ("Fields", "fields"), ("Doubted fields", "doubted"), ("Peer comments", "peer_comments"), ("Peer mentions", "peer_mentions"), ("Checked at source", "checked")], data["sources"])
    sheet("Peer comments", [("Project", "project"), ("Source", "source"), ("Record", "record"), ("Field", "field"), ("Reviewer", "by"), ("Role", "role"), ("Kind", "type"), ("Comment", "comment"), ("When", "when")], data["comments"])
    sheet("Fields", [("Project", "project"), ("Source", "source"), ("Website", "website"), ("Record", "record"), ("Field", "field"), ("Value", "value"), ("Doubted", "doubted"), ("Why doubted", "why"),
                     ("Peer says", "peer_says"), ("Peer mentioned in", "peer_mentioned_in"), ("Other flags", "other"), ("Admin verdict", "verdict"), ("Real value", "truth")], data["fields"])
    sheet("Records JSON", [("Project", "project"), ("Source", "source"), ("Record", "record"), ("Record JSON", "json")], data["records"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def extractor_xlsx(report: dict, error_rows: list[dict], meta: dict) -> bytes:
    """The extractor assessment: one row per extractor, one per source, and every error with what happened to it."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    for i, line in enumerate([
        "Extractor accuracy and error rate.",
        "",
        "Every error stays in these numbers after the extractor fixes it. An error is one field on one record, counted once.",
        "Errors per 100 fields = errors / fields filled in x 100.",
        "First-pass accuracy = checked fields that were right the first time an admin checked them.",
        "Final accuracy = checked fields that are right now. Checks on values the extractor changed since are waiting for a re-check and are left out.",
        "Caught by reviewers = errors a reviewer pointed at before the admin did. Fixed = errors that are right now.",
        "Rework rate = records that were sent back at least once / records.",
        "Under 30 checked fields the numbers are low confidence.",
        "",
        "Period: %s. Generated %s." % (meta.get("period", "all time"), meta.get("generated", "")),
    ], 1):
        ws.cell(i, 1, line)
    ws.column_dimensions["A"].width = 120
    head = PatternFill("solid", fgColor="1F2937")

    def sheet(title, cols, rows):
        w = wb.create_sheet(title)
        w.append(cols)
        for c in w[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head
        for r in rows:
            w.append([_cell(v) for v in r])
        w.freeze_panes = "A2"
        for i, c in enumerate(cols, 1):
            w.column_dimensions[get_column_letter(i)].width = max(12, min(48, len(c) + 4))
        return w

    ex = report["extractors"]
    sheet("Extractors", ["Extractor", "Sources", "Records", "Fields", "Errors", "Errors per 100 fields", "Checked fields", "First-pass accuracy %", "Final accuracy %", "Waiting re-check",
                         "Caught by reviewers", "Fixed", "Still wrong", "Records sent back", "Rework rate %", "Times sent back", "Confidence"],
          [[x["name"], len(x["sources"]), x["records"], x["fields"], x["errors"], x["errors_per_100"], x["audited"], x["first_pass_accuracy"], x["final_accuracy"], x["waiting_recheck"],
            x["caught"], x["fixed"], x["still_wrong"], x["returned_records"], x["rework_rate"], x["returns"], x["confidence"]] for x in ex])
    sheet("Per source", ["Extractor", "Source", "Records", "Fields", "Errors", "Errors per 100 fields", "Checked fields", "First-pass accuracy %", "Final accuracy %", "Waiting re-check",
                         "Caught by reviewers", "Fixed", "Still wrong", "Records sent back", "Times sent back"],
          [[x["name"], d["source"], d["records"], d["fields"], d["errors"], d["errors_per_100"], d["audited"], d["first_pass_accuracy"], d["final_accuracy"], d["waiting_recheck"],
            d["caught"], d["fixed"], d["still_wrong"], d["returned_records"], d["returns"]] for x in ex for d in x["sources"]])
    sheet("Errors", ["Extractor", "Source", "Record", "Field", "Value when found", "Value now", "What the source says", "Found by", "Caught by reviewer", "Status", "Kind", "Severity", "Note"],
          [[e["extractor"], e["source"], e["record"], e["field"], e["was"], e["now"], e["truth"], e["found_by"], e["caught_by"], e["status"], e["kind"], e["severity"], e["note"]] for e in error_rows])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def sop_requirements_xlsx(items: list[dict], meta: dict) -> bytes:
    """What each source must follow, straight from its saved SOP: one row per rule, field and warning, each with its section and quote."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    for i, line in enumerate([
        "SOP requirements per source.",
        "",
        "Everything here is read from the SOP saved on each source. Nothing is added or guessed.",
        "Sources: one row per source with its category, SOP code and version, and how many rules and warnings it has.",
        "Fields: every field the SOP names, whether it is required, its allowed values and format, the section it came from and the exact line.",
        "Rules: every sentence in the SOP that says must, shall, never, only, exactly, at least and similar. 'Matches this source' is yes when the section names the source's category or entity.",
        "Warnings: things to settle before judging a source strictly: no SOP saved, no text, a different SOP code in the text, no section for the category, a missing version.",
        "MUST rules are strict. SHOULD rules are advice. A rule under another category's heading is shown but not enforced on this source.",
        "",
        "Generated %s." % meta.get("generated", ""),
    ], 1):
        ws.cell(i, 1, line)
    ws.column_dimensions["A"].width = 120
    head = PatternFill("solid", fgColor="1F2937")

    def sheet(title, cols, rows):
        w = wb.create_sheet(title)
        w.append(cols)
        for c in w[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head
        for r in rows:
            w.append([_cell(v) for v in r])
        w.freeze_panes = "A2"
        for i, c in enumerate(cols, 1):
            w.column_dimensions[get_column_letter(i)].width = max(14, min(70, len(c) + 6)) if c not in ("Rule", "Quote", "Warning") else 90
            if c in ("Rule", "Quote", "Warning"):
                for cell in w[get_column_letter(i)][1:]:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        return w
    sheet("Sources", ["Source", "Category", "SOP code", "SOP version", "Entity", "SOP title", "Has SOP text", "Required fields", "Fields", "Sections", "Sections naming this source", "Rules", "MUST rules", "Warnings"],
          [[x["name"], x["category"], x["sop_code"], x["spec_version"], x["entity"], x["title"], "yes" if x["has_markdown"] else "no", x["counts"].get("required_fields", len(x["required_fields"])), x["counts"].get("fields", ""),
            x["counts"].get("sections", ""), x["counts"].get("matching_sections", ""), x["counts"].get("rules", ""), x["counts"].get("must_rules", ""), len(x["warnings"])] for x in items])
    sheet("Fields", ["Source", "Category", "SOP code", "Field", "Required", "Type", "Allowed values", "Format", "Rule", "Section", "Enforced on this source", "Quote"],
          [[x["name"], x["category"], x["sop_code"], f["name"], "yes" if f["required"] else ("no" if f["required"] is False else "not stated"), f["type"], ", ".join(f["allowed"] or []), f["format"], f["rule"], f["section"],
            "yes" if f.get("enforce", True) else "no (another category's section)", f["quote"]] for x in items for f in x["fields"]])
    sheet("Rules", ["Source", "Category", "SOP code", "Section", "Matches this source", "Strength", "Kind", "Fields named", "Rule"],
          [[x["name"], x["category"], x["sop_code"], " > ".join(sx["path"]) or sx["title"], "yes" if sx["matches"] else "no", r["strength"], r["kind"], ", ".join(r["fields"] or []), r["text"]] for x in items for sx in x["sections"] for r in sx["rules"]])
    sheet("Warnings", ["Source", "Category", "SOP code", "Warning"], [[x["name"], x["category"], x["sop_code"], w] for x in items for w in x["warnings"]])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def process_workbook(sc: dict, meta: dict) -> bytes:
    """The whole process in one workbook: the cumulative total, one row per source (batch), each extractor and each peer reviewer.
    Every number comes from the same counts, and the README says exactly how each rate is worked out."""
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    window = " to ".join(x for x in (meta.get("date_from"), meta.get("date_to")) if x) or "everything ever recorded"
    gen = _fmt(meta.get("generated"))
    lines = [
        "Process report: accuracy, errors and review catch across the sources.",
        f"Period: {window}.  Generated: {gen} (UTC).",
        "",
        "Sheets",
        "Cumulative: the total across all sources in this report.",
        "Pulls: one row per batch, which is the set of sources pulled from Xtrium on the same day (Nairobi time), oldest first. The last columns show the running total up to and including that batch.",
        "Batch extractors: each extractor's figures inside each batch.",
        "Batch reviewers: each peer reviewer's catch rate inside each batch, on the sources they were assigned in it.",
        "Sources: one row per source, highest error rate first.",
        "Extractors: the same figures across all batches, grouped by the extractor assigned to each source.",
        "Reviewers: each peer reviewer's catches, false alarms and catch rate across all batches.",
        "",
        "How each number is worked out",
        "Error: one record and field that an admin confirmed as wrong, or that an admin's check against the real source found wrong. Counted once, however many people pointed at it. Errors stay counted after the extractor fixes them.",
        "Errors per 100 fields = errors / fields delivered x 100.",
        "First-pass accuracy = checked fields that were right the first time / checked fields. It never changes after a fix.",
        "Final accuracy = checked fields that are right now / checked fields whose current value was checked. A field changed since its check is 'waiting re-check' and left out until it is checked again.",
        "Catch rate = errors a peer reviewer pointed at / all known errors. Only errors an admin has confirmed or found count, so it can rise as more fields are checked.",
        "Reviewer catch rate = errors that reviewer pointed at / all known errors on the sources they were assigned to.",
        "Reviewer precision = fields the reviewer flagged that were confirmed / (confirmed + dismissed).",
        "Rework rate = records sent back at least once / records.",
        "Confidence is 'low' when fewer than 30 fields were checked (fewer than 10 known errors for a reviewer). Treat low-confidence rates as early signals.",
        "Edits made after peers began reviewing are locked. An admin override is logged, so scores can be re-run with it excluded.",
    ]
    for i, line in enumerate(lines, 1):
        ws.cell(i, 1, line)
        if i in (1, 4, 13):
            ws.cell(i, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 150
    head = PatternFill("solid", fgColor="1F2937")

    def sheet(title, cols, rows):
        w = wb.create_sheet(title)
        w.append(cols)
        for c in w[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head
        for r in rows:
            w.append([_cell(v) for v in r])
        w.freeze_panes = "A2"
        for i, c in enumerate(cols, 1):
            w.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(c) + 4))
        return w

    def pc(v):
        return "" if v is None else v / 100.0
    num_cols = ["Records", "Fields", "Checked", "Errors", "Errors per 100 fields", "First-pass accuracy", "Final accuracy", "Caught by peers", "Missed by peers", "Catch rate", "Fixed", "Still wrong", "Records sent back", "Rework rate", "Confidence"]

    def nums(d):
        return [d["records"], d["fields"], d["audited"], d["errors"], d["errors_per_100"], pc(d["first_pass_accuracy"]), pc(d["final_accuracy"]), d["caught"], d["escaped"], pc(d["caught_rate"]),
                d["fixed"], d["still_wrong"], d["returned_records"], pc(d["rework_rate"]), d["confidence"]]
    c = sc["cumulative"]
    sheet("Cumulative", ["Sources"] + num_cols, [[c.get("sources", 0)] + nums(c)])
    run_cols = ["Sources so far", "Fields so far", "Errors so far", "Errors per 100 so far", "First-pass accuracy so far", "Final accuracy so far", "Catch rate so far", "Rework rate so far"]
    pulls = sc.get("pulls") or []
    if pulls:
        sheet("Pulls", ["Batch", "Pulled on", "Sources"] + num_cols + run_cols,
              [[b["label"], b["date"], b["sources"]] + nums(b) + [b["cumulative"]["sources"], b["cumulative"]["fields"], b["cumulative"]["errors"], b["cumulative"]["errors_per_100"], pc(b["cumulative"]["first_pass_accuracy"]),
                pc(b["cumulative"]["final_accuracy"]), pc(b["cumulative"]["caught_rate"]), pc(b["cumulative"]["rework_rate"])] for b in pulls])
        sheet("Batch extractors", ["Batch", "Extractor", "Sources"] + num_cols, [[b["label"], x["name"], x["sources"]] + nums(x) for b in pulls for x in b["extractors"]])
        sheet("Batch reviewers", ["Batch", "Reviewer", "Sources assigned", "Known errors on them", "Caught", "Missed", "Catch rate", "False alarms", "Precision", "Only one to catch"],
              [[b["label"], r["name"], r["sources"], r["known_errors"], r["caught"], r["missed"], pc(r["catch_rate"]), r["false_alarms"], pc(r["precision"]), r["only_one_to_catch"]] for b in pulls for r in b["reviewers"]])
    sheet("Sources", ["Source", "Status", "Extractor", "Peer reviewers"] + num_cols,
          [[b["source"], b.get("status", ""), b["extractor"], ", ".join(b["peers"])] + nums(b) for b in sc["batches"]])
    sheet("Extractors", ["Extractor", "Sources"] + num_cols, [[x["name"], len(x["sources"])] + nums(x) for x in sc["extractors"]])
    sheet("Reviewers", ["Reviewer", "Sources assigned", "Known errors on them", "Caught", "Missed", "Catch rate", "False alarms", "Precision", "Still open", "Only one to catch", "Confidence"],
          [[r["name"], r["sources"], r["known_errors"], r["caught"], r["missed"], pc(r["catch_rate"]), r["false_alarms"], pc(r["precision"]), r["open"], r["only_one_to_catch"], r["confidence"]] for r in sc["reviewers"]])
    for name in [n for n in wb.sheetnames if n != "README"]:
        w = wb[name]
        for row in w.iter_rows(min_row=2):
            for cell in row:
                h = w.cell(1, cell.column).value
                if h and (str(h).endswith("accuracy") or str(h).endswith("rate") or h in ("Precision",) or "accuracy so far" in str(h) or "rate so far" in str(h)) and isinstance(cell.value, float):
                    cell.number_format = "0.0%"
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def final_check_xlsx(res: dict, meta: dict) -> bytes:
    """The final-records check: summary, per source, every wrong field with who flagged it, false alarms, extractors and reviewers."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    for i, line in enumerate([
        "Final records checked against what the extractor delivered and what the reviewers flagged.",
        "",
        "Error % = extractor errors / fields compared x 100. Accuracy % = 100 - error %. Held fields and admin changes are in the fields compared but are not errors.",
        "Extractor error = the delivered value was wrong, unsupported or missing, or you marked it wrong in review (even when the final value is still the delivered one: Not fixed).",
        "Admin change = the extractor's value was right and the final value is your rewording, reordering or optional addition. It is not counted against the extractor.",
        "Held = it cannot be decided from what is recorded. Decide it in the Final check, or record a verdict in Check against source.",
        "A delivered record with no final record counts as wrong, every field of it, unless you overruled it. SOP failures are checked on the final records and are not extractor errors.",
        "A field is compared when either side has a value. Timestamps (%s) are not compared." % (", ".join(res.get("ignored_fields") or []) or "none"),
        "Wording that differs only in case, spacing or punctuation counts as correct.",
        "Caught = a wrong field a reviewer had flagged. Got through = a wrong field nobody flagged. False alarm = a field a reviewer flagged that was right.",
        "Reviewer catch rate = wrong fields they flagged / all wrong fields on their sources. Precision = their flags that were really wrong.",
        "To confirm = the final value differs and the admin has not decided yet. Held fields are not counted as errors and nothing is recorded for them.",
        "Admin change = the admin decided the extractor was right and the final value is a change. Counted correct.",
        "Fixed after review = a reviewer flagged an earlier value that the extractor then fixed. Counted as a reviewer catch, not as an error.",
        "Nothing was edited. The extractor's values are as they were delivered.",
        "",
        "Generated %s. %s" % (meta.get("generated", ""), "Verdicts were recorded as audits." if res.get("mode") == "applied" else "Check only: nothing was recorded."),
    ], 1):
        ws.cell(i, 1, line)
    ws.column_dimensions["A"].width = 120
    head = PatternFill("solid", fgColor="1F2937")

    def sheet(title, cols, rows, widths=None):
        w = wb.create_sheet(title)
        w.append(cols)
        for c in w[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head
        for r in rows:
            w.append([_cell(v) for v in r])
        w.freeze_panes = "A2"
        for i, c in enumerate(cols, 1):
            w.column_dimensions[get_column_letter(i)].width = (widths or {}).get(c, max(12, min(48, len(c) + 4)))
        for row in w.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        return w

    T = res["totals"]
    sheet("Summary", ["Measure", "Value"], [
        ["Accuracy %", T["accuracy_pct"]], ["Error rate %", T["error_pct"]], ["Fields compared", T["fields"]], ["Correct", T.get("correct", 0)], ["Wrong fields", T["wrong"]],
        ["Of which not fixed", T.get("not_fixed", 0)], ["Of which in records with no final match", T.get("unmatched_fields", 0)], ["SOP failures (final records)", T.get("sop_failures", 0)],
        ["Caught by reviewers", T["caught"]], ["Fixed after review (catches)", T.get("fixed", 0)], ["Got through", T["escaped"]], ["To confirm (held)", T.get("pending", 0)], ["Admin changes (extractor right)", T.get("admin_changes", 0)], ["Reviewer catch rate %", T["catch_rate"]], ["False alarms", T["false_alarms"]],
        ["Wording-only differences (counted correct)", T["cosmetic"]], ["Sources", T["sources"]], ["Records matched", T["records"]],
        ["Check: wrong + admin changes + held + correct = fields", "%s = %s" % (T["wrong"] + T["admin_changes"] + T["pending"] + T.get("correct", 0), T["fields"])]], {"Measure": 52})
    sheet("By source", ["Source", "Extractor", "Reviewers", "Fields", "Extractor errors", "Not fixed", "Admin changes", "Held", "SOP failures", "Error %", "Accuracy %", "Caught", "Fixed after review", "Got through", "False alarms", "Catch rate %"],
          [[s["name"], s["extractor"], ", ".join(s["reviewers"]), s["fields"], s["wrong"], s.get("not_fixed", 0), s.get("admin_changes", 0), s.get("pending", 0), s.get("sop_failures", 0), s["error_pct"], s["accuracy_pct"], s["caught"], s.get("fixed", 0), s["escaped"], s["false_alarms"], s["catch_rate"]] for s in res["sources"]],
          {"Source": 36, "Reviewers": 30})
    sheet("Wrong fields", ["Source", "Extractor", "Record", "Field", "Extractor delivered", "Final value", "Kind", "Class", "Reason", "Flagged by", "System check flagged", "Result"],
          [[s["name"], s["extractor"], w["record"], w["field"], w["extractor_value"], w["final_value"], w["category"],
            {"extractor_error": "Extractor error", "admin_change": "Admin change", "held": "Held"}.get(w.get("class"), ""), w.get("reason", ""), ", ".join(w["flagged_by"]), "yes" if w["system_flagged"] else "",
            {"caught": "Caught", "escaped": "Got through", "to_confirm": "To confirm", "admin_change": "Admin change"}.get(w["outcome"], w["outcome"])] for s in res["sources"] for w in s["wrong_fields"]],
          {"Source": 30, "Extractor delivered": 44, "Final value": 44, "Field": 26, "Reason": 56})
    sheet("Not fixed", ["Source", "Extractor", "Record", "Field", "Value (still the delivered value)", "Flagged by", "Result"],
          [[s["name"], s["extractor"], w["record"], w["field"], w["extractor_value"], ", ".join(w["flagged_by"]), "Caught" if w["outcome"] == "caught" else "Got through"]
           for s in res["sources"] for w in s["wrong_fields"] if w.get("not_fixed")], {"Source": 30, "Value (still the delivered value)": 50})
    sheet("Held", ["Source", "Record", "Field", "Extractor delivered", "Final value", "What is needed"],
          [[s["name"], w["record"], w["field"], w["extractor_value"], w["final_value"], w.get("reason", "")] for s in res["sources"] for w in s["wrong_fields"] if w.get("class") == "held"],
          {"Source": 30, "Extractor delivered": 44, "Final value": 44, "What is needed": 60})
    sheet("Records with no match", ["Source", "Delivered record", "Fields scored", "Counted as", "Final record that replaced it"],
          [[s["name"], u["record"], u["fields"], "Wrong" if u["decision"] == "wrong" else "Not an error (you overruled it)", ", ".join(u["replaced_by"])] for s in res["sources"] for u in s.get("unmatched_scored", [])],
          {"Source": 30, "Delivered record": 36, "Final record that replaced it": 40})
    sheet("SOP failures", ["Source", "Extractor", "Record", "Field", "Rule", "Detail"],
          [[s["name"], s["extractor"], f["record"], f["field"], f["rule"], f["why"]] for s in res["sources"] for f in s.get("sop_failure_list", [])], {"Source": 30, "Detail": 70})
    sheet("Fixed after review", ["Source", "Record", "Field", "Flagged value", "Now", "Flagged by"],
          [[s["name"], f["record"], f["field"], f["was"], f["now"], f["flagged_by"]] for s in res["sources"] for f in s.get("fixed_fields", [])], {"Source": 30, "Flagged value": 40, "Now": 40})
    sheet("False alarms", ["Source", "Record", "Field", "Value", "Flagged by"],
          [[s["name"], f["record"], f["field"], f["value"], ", ".join(f["flagged_by"])] for s in res["sources"] for f in s["false_alarm_fields"]], {"Source": 30, "Value": 44})
    sheet("Extractors", ["Extractor", "Sources", "Fields", "Extractor errors", "Not fixed", "Admin changes", "Held", "SOP failures", "Error %", "Accuracy %"], [[e["name"], e["sources"], e["fields"], e["wrong"], e.get("not_fixed", 0), e.get("admin_changes", 0), e.get("pending", 0), e.get("sop_failures", 0), e["error_pct"], e["accuracy_pct"]] for e in res["extractors"]], {"Extractor": 28})
    sheet("Reviewers", ["Reviewer", "Sources", "Flags", "Caught", "Missed", "False alarms", "Wrong fields on their sources", "Catch rate %", "Precision %"],
          [[r["name"], r["sources"], r["flags"], r["caught"], r.get("missed", 0), r["false_alarms"], r["errors_on_their_sources"], r["catch_rate"], r["precision"]] for r in res["reviewers"]], {"Reviewer": 28})
    rows = [[s["name"], "Final record with no match", ", ".join(s["key_records_unmatched"])] for s in res["sources"] if s["key_records_unmatched"]]
    rows += [[s["name"], "Extractor record with no match", ", ".join(s["extractor_records_unmatched"])] for s in res["sources"] if s["extractor_records_unmatched"]]
    rows += [[n, "Source not found", ""] for n in res["unmatched_sources"]]
    sheet("Not scored", ["Source", "Why", "Records"], rows, {"Source": 36, "Why": 32, "Records": 50})
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def rounds_xlsx(res: dict, meta: dict) -> bytes:
    """What changed between the kept first round (v1) and the records now (v2)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    for i, line in enumerate([
        "v1 is the first extraction, kept when the source was re-extracted. v2 is what the records say now.",
        "Records are matched by name. Fields are compared in both directions. Timestamps (%s) are ignored." % (", ".join(res.get("ignored_fields") or []) or "none"),
        "Same = identical. Wording = only case, spacing or punctuation. Changed = both rounds have a value and they differ. Added = v2 has a value v1 did not. Removed = v1 had a value v2 does not.",
        "Fixed = v1 was checked wrong and v2 is right. Regressed = v1 was right and v2 is wrong. Still wrong = wrong in both. Unchecked = no verdict to say which is right.",
        "",
        "Generated %s." % meta.get("generated", ""),
    ], 1):
        ws.cell(i, 1, line)
    ws.column_dimensions["A"].width = 130
    head = PatternFill("solid", fgColor="1F2937")

    def sheet(title, cols, rows, widths=None):
        w = wb.create_sheet(title)
        w.append(cols)
        for c in w[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head
        for r in rows:
            w.append([_cell(v) for v in r])
        w.freeze_panes = "A2"
        for i, c in enumerate(cols, 1):
            w.column_dimensions[get_column_letter(i)].width = (widths or {}).get(c, max(12, min(48, len(c) + 4)))
        for row in w.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")

    T = res["totals"]
    sheet("Summary", ["Measure", "Value"], [
        ["Sources compared", T["sources"]], ["v1 records", T["v1_records"]], ["v2 records", T["v2_records"]], ["Fields compared", T["fields"]],
        ["Same", T["same"]], ["Wording only", T["wording"]], ["Changed", T["changed"]], ["Added in v2", T["added"]], ["Removed in v2", T["removed"]], ["Changed %", T["change_pct"]],
        ["Fixed (v1 wrong, v2 right)", T["fixed"]], ["Regressed (v1 right, v2 wrong)", T["regressed"]], ["Still wrong", T["still_wrong"]], ["Unchecked", T["unchecked"]],
        ["v1 fields checked", T["v1_checked"]], ["v1 wrong", T["v1_wrong"]], ["v1 error %", T["v1_error_pct"]],
        ["v2 fields checked", T["v2_checked"]], ["v2 wrong", T["v2_wrong"]], ["v2 error %", T["v2_error_pct"]],
        ["v1 records dropped", T["records_dropped"]], ["New records in v2", T["records_new"]]], {"Measure": 36})
    sheet("By source", ["Source", "v1 extractor", "v2 extractor", "v1 reviewers", "v2 reviewers", "v1 records", "v2 records", "Fields", "Same", "Wording", "Changed", "Added", "Removed", "Changed %", "Fixed", "Regressed", "Still wrong", "Unchecked"],
          [[s["name"], s["v1_extractor"], s["v2_extractor"], ", ".join(s["v1_reviewers"]), ", ".join(s["v2_reviewers"]), s["v1_records"], s["v2_records"], s["fields"], s["same"], s["wording"], s["changed"],
            s["added"], s["removed"], s["change_pct"], s["fixed"], s["regressed"], s["still_wrong"], s["unchecked"]] for s in res["sources"]], {"Source": 34, "v1 reviewers": 24, "v2 reviewers": 24})
    sheet("Changed fields", ["Source", "Record", "Field", "Kind", "v1 value", "v2 value", "Result", "v1 extractor", "v2 extractor"],
          [[s["name"], c["record"], c["field"], c["kind"], c["v1_value"], c["v2_value"], {"fixed": "Fixed", "regressed": "Regressed", "still_wrong": "Still wrong"}.get(c["result"], "Unchecked"),
            s["v1_extractor"], s["v2_extractor"]] for s in res["sources"] for c in s["changes"]], {"Source": 30, "v1 value": 44, "v2 value": 44, "Field": 26})
    sheet("Records dropped or new", ["Source", "What", "Records"],
          [[s["name"], "In v1 only (dropped in v2)", ", ".join(s["dropped_records"])] for s in res["sources"] if s["dropped_records"]] +
          [[s["name"], "In v2 only (new)", ", ".join(s["new_records"])] for s in res["sources"] if s["new_records"]], {"Source": 34, "What": 28, "Records": 60})
    sheet("People", ["Person", "Sources in v1", "Sources in v2"], [[p["name"], p["v1_sources"], p["v2_sources"]] for p in res["people"]], {"Person": 28})
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
