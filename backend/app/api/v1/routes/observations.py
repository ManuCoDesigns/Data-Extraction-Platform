"""Observations: what extractors and reviewers saw on each source (not accessible, needs a login, no data...).
Each extractor and peer reviewer has to record theirs, even if it is "no issues found"."""
import io

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.api.v1.routes.enhancements import _can_chat, _need_lead, _visible_sources
from app.core.security import get_current_user
from app.db.session import get_db
from app.models.all_models import Source, SourceLinkCheck, SourceObservation, SourceStatus, User
from app.services import enh, observations as ob

router = APIRouter(prefix="/observations", tags=["observations"])


class ObsItem(BaseModel):
    kind: str
    note: str | None = None
    record_id: str | None = None


class ObsMine(BaseModel):
    none: bool = False
    items: list[ObsItem] = []


def _src(db: Session, user: User, source_id: str) -> Source:
    s = _visible_sources(db, user).filter(Source.id == source_id).first()
    if not s or not _can_chat(user, s):
        raise HTTPException(status_code=404, detail="Source not found.")
    return s


def _names(db: Session, ids: set) -> dict:
    ids = {i for i in ids if i}
    return {u.id: u.full_name for u in db.query(User).filter(User.id.in_(ids)).all()} if ids else {}


def _row(o: SourceObservation, names: dict) -> dict:
    return {"id": o.id, "source_id": o.source_id, "user_id": o.user_id, "user": names.get(o.user_id), "role": o.role, "kind": o.kind,
            "label": "No issues found" if o.kind == "none" else ob.LABEL.get(o.kind, o.kind), "note": o.note, "record_id": o.record_id, "created_at": o.created_at}


@router.get("/kinds")
def kinds(user: User = Depends(get_current_user)):
    return {"kinds": ob.kind_list()}


@router.get("/sources/{source_id}")
def for_source(source_id: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    s = _src(db, user, source_id)
    rows = ob.rows_for(db, [s.id])
    owes = ob.who_owes(s, rows)
    names = _names(db, {r.user_id for r in rows} | {o["user_id"] for o in owes})
    mine_role = ob.role_on(s, user)
    mine = [_row(r, names) for r in rows if r.user_id == user.id]
    return {
        "source_id": s.id, "my_role": mine_role, "i_owe": bool(mine_role in ("extractor", "reviewer") and not mine),
        "mine": mine, "none": bool(mine and mine[0]["kind"] == "none"),
        "others": [_row(r, names) for r in rows if r.user_id != user.id],
        "people": [{**o, "name": names.get(o["user_id"])} for o in owes],
        "kinds": ob.kind_list(),
    }


@router.put("/sources/{source_id}/mine")
def save_mine(source_id: str, body: ObsMine, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    s = _src(db, user, source_id)
    role = ob.role_on(s, user)
    if not role:
        raise HTTPException(status_code=403, detail="Only the extractor, the peer reviewers and admins record observations on a source.")
    ob.replace_mine(db, s, user, role, body.none, [i.model_dump() for i in body.items])
    db.commit()
    return for_source(source_id, db, user)


@router.get("/pending")
def pending(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """Sources where this person still has to record their observations."""
    srcs = [s for s in _visible_sources(db, user).filter(Source.status != SourceStatus.APPROVED).all()
            if s.assigned_extractor_id == user.id or any(p.user_id == user.id for p in s.peer_reviews)]
    rows = ob.rows_for(db, [s.id for s in srcs])
    done = {r.source_id for r in rows if r.user_id == user.id}
    items = [{"source_id": s.id, "project_id": s.project_id, "name": s.name, "role": "extractor" if s.assigned_extractor_id == user.id else "reviewer",
              "status": getattr(s.status, "value", s.status)}
             for s in srcs if s.id not in done and (s.assigned_extractor_id == user.id or s.status not in (SourceStatus.NOT_STARTED, SourceStatus.EXTRACTING))]
    return {"count": len(items), "items": items}


def _collect(db: Session, user: User, project_id, kind, source_id, user_id):
    srcs = _visible_sources(db, user, project_id).all()
    if source_id:
        srcs = [s for s in srcs if s.id == source_id]
    by = {s.id: s for s in srcs}
    rows = ob.rows_for(db, list(by))
    names = _names(db, {r.user_id for r in rows} | {s.assigned_extractor_id for s in srcs} | {p.user_id for s in srcs for p in s.peer_reviews})
    sel = [r for r in rows if r.kind != "none" and (not kind or r.kind == kind) and (not user_id or r.user_id == user_id)]
    missing = []
    for s in srcs:
        if s.status in (SourceStatus.APPROVED, SourceStatus.NOT_STARTED):
            continue
        for o in ob.who_owes(s, rows):
            if not o["done"] and (o["role"] == "extractor" or s.status != SourceStatus.EXTRACTING):
                missing.append({"source_id": s.id, "project_id": s.project_id, "source": s.name, "user_id": o["user_id"], "person": names.get(o["user_id"]), "role": o["role"]})
    return srcs, by, rows, sel, missing, names


@router.get("")
def summary(project_id: str | None = None, kind: str | None = None, source_id: str | None = None, user_id: str | None = None, export: str | None = None,
            db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """Everything recorded, how often each kind came up, and who still owes a check. Admins and QA leads."""
    _need_lead(user)
    srcs, by, rows, sel, missing, names = _collect(db, user, project_id, kind, source_id, user_id)
    items = [{**_row(r, names), "source": by[r.source_id].name, "project_id": by[r.source_id].project_id, "item_no": by[r.source_id].external_ref_id,
              "group": ob.KINDS.get(r.kind, ("", "Content", False))[1]} for r in sorted(sel, key=lambda r: r.created_at, reverse=True)]
    if export == "xlsx":
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Observations"
        ws.append(["Source", "Xtrium item", "Issue", "Group", "Recorded by", "Role", "Note", "Recorded"])
        for i in items:
            ws.append([i["source"], i["item_no"], i["label"], i["group"], i["user"], i["role"], i["note"], str(i["created_at"] or "")[:16]])
        w2 = wb.create_sheet("By kind")
        w2.append(["Issue", "Times recorded", "Sources"])
        for c in ob.counts(sel):
            w2.append([c["label"], c["count"], c["sources"]])
        w3 = wb.create_sheet("Still to record")
        w3.append(["Source", "Person", "Role"])
        for m in missing:
            w3.append([m["source"], m["person"], m["role"]])
        buf = io.BytesIO()
        wb.save(buf)
        return Response(content=buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": 'attachment; filename="observations.xlsx"'})
    return {
        "generated_at": enh.now_utc(), "total": len(items), "sources_affected": len({i["source_id"] for i in items}),
        "by_kind": ob.counts(sel), "items": items[:500], "missing": missing,
        "confirmed_none": sum(1 for r in rows if r.kind == "none"), "kinds": ob.kind_list(),
    }


# ─── Link triage: flag the sources that will not open, at the very start ──────────

class TriageItem(BaseModel):
    source_id: str
    kind: str
    note: str | None = None


class TriageBody(BaseModel):
    items: list[TriageItem]


@router.get("/triage")
def triage_list(project_id: str | None = None, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """The sources this person works, with the last link check and what they have already flagged.
    Extractors see their own; admins see every source that is not yet approved."""
    admin = ob._is_admin(user)
    srcs = [s for s in _visible_sources(db, user, project_id).filter(Source.status != SourceStatus.APPROVED).order_by(Source.name).all()
            if s.assigned_extractor_id == user.id or admin]
    ids = [s.id for s in srcs]
    checks = {c.source_id: c for c in db.query(SourceLinkCheck).filter(SourceLinkCheck.source_id.in_(ids)).all()} if ids else {}
    rows = ob.rows_for(db, ids)
    items = []
    for s in srcs:
        c = checks.get(s.id)
        stale = bool(c and c.url and c.url != s.website_url)
        res = None if (not c or stale) else c.status
        flags = [{"kind": r.kind, "label": "No issues found" if r.kind == "none" else ob.LABEL.get(r.kind, r.kind), "note": r.note, "mine": r.user_id == user.id}
                 for r in rows if r.source_id == s.id and r.kind != "none"]
        items.append({"source_id": s.id, "project_id": s.project_id, "name": s.name, "url": s.website_url, "item_no": s.external_ref_id,
                      "status": getattr(s.status, "value", s.status), "mine": s.assigned_extractor_id == user.id,
                      "link": {"status": res or "unchecked", "detail": c.detail if c and not stale else None, "suggest": ob.FROM_LINK.get(res or "")},
                      "flags": flags, "flagged": any(f["mine"] for f in flags)})
    return {"items": items, "total": len(items), "flagged": sum(1 for i in items if i["flags"]), "kinds": ob.kind_list()}


@router.post("/triage")
def triage_save(body: TriageBody, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    if not body.items:
        raise HTTPException(status_code=422, detail="Nothing to save.")
    ids = [i.source_id for i in body.items]
    by = {s.id: s for s in _visible_sources(db, user).filter(Source.id.in_(ids)).all()}
    done, skipped = 0, []
    for it in body.items:
        s = by.get(it.source_id)
        role = ob.role_on(s, user) if s else None
        if not s or not role:
            skipped.append({"source_id": it.source_id, "why": "You do not work this source."}); continue
        ob.add_mine(db, s, user, role, it.kind, it.note)
        done += 1
    db.commit()
    return {"saved": done, "skipped": skipped}
